from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook


PREVIEW_ROWS = 8


def _json_cell(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _normalize_sheet_name(value: Any) -> str:
    return str(value or "").replace("\u00a0", " ").strip().casefold()


def _dedupe_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, raw in enumerate(values, start=1):
        base = str(raw or "").strip() or f"Column_{index}"
        count = counts.get(base, 0)
        counts[base] = count + 1
        headers.append(base if count == 0 else f"{base}.{count}")
    return headers


def _resolve_sheet_name(sheet_names: list[str], requested: str | None) -> str:
    if not sheet_names:
        raise HTTPException(status_code=400, detail="В Excel не найдено листов.")

    if requested:
        if requested in sheet_names:
            return requested
        normalized = _normalize_sheet_name(requested)
        match = next(
            (name for name in sheet_names if _normalize_sheet_name(name) == normalized),
            None,
        )
        if match:
            return match

    return sheet_names[0]


def _preview_xlsx(
    file_bytes: bytes,
    *,
    requested_sheet: str | None,
    header_row: int,
) -> dict[str, Any]:
    try:
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Не удалось открыть Excel на сервере: {exc}",
        ) from exc

    try:
        sheet_names = list(workbook.sheetnames)
        selected_sheet = _resolve_sheet_name(sheet_names, requested_sheet)
        worksheet = workbook[selected_sheet]

        header_values = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            ),
            (),
        )
        columns = _dedupe_headers(list(header_values))

        rows: list[dict[str, Any]] = []
        for values in worksheet.iter_rows(
            min_row=header_row + 1,
            max_row=header_row + PREVIEW_ROWS,
            values_only=True,
        ):
            if not columns:
                break
            row = {
                column: _json_cell(values[index] if index < len(values) else None)
                for index, column in enumerate(columns)
            }
            rows.append(row)

        max_row = int(worksheet.max_row or 0)
        row_count = max(0, max_row - header_row)
        return {
            "rows": rows,
            "columns": columns,
            "sheetNames": sheet_names,
            "selectedSheet": selected_sheet,
            "headerRow": header_row,
            "previewRows": rows,
            "rowCount": row_count,
            "parser": "server-openpyxl",
        }
    finally:
        workbook.close()


def _preview_csv(file_bytes: bytes, *, header_row: int) -> dict[str, Any]:
    decoded: str | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            decoded = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise HTTPException(status_code=400, detail="Не удалось определить кодировку CSV.")

    lines = decoded.splitlines()
    if len(lines) < header_row:
        raise HTTPException(
            status_code=400,
            detail=f"В CSV нет строки заголовков №{header_row}.",
        )

    sample = "\n".join(lines[max(0, header_row - 1): header_row + 15])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"

    reader = csv.reader(lines[header_row - 1:], delimiter=delimiter)
    header_values = next(reader, [])
    columns = _dedupe_headers(header_values)

    rows: list[dict[str, Any]] = []
    for values in reader:
        rows.append({
            column: values[index] if index < len(values) else ""
            for index, column in enumerate(columns)
        })
        if len(rows) >= PREVIEW_ROWS:
            break

    row_count = max(0, len(lines) - header_row)
    return {
        "rows": rows,
        "columns": columns,
        "sheetNames": [],
        "selectedSheet": "",
        "headerRow": header_row,
        "previewRows": rows,
        "rowCount": row_count,
        "parser": "server-csv",
    }


def preview_source_file(
    file_bytes: bytes,
    file_name: str,
    *,
    requested_sheet: str | None = None,
    header_row: int = 1,
) -> dict[str, Any]:
    try:
        normalized_header_row = max(1, min(200, int(header_row or 1)))
    except (TypeError, ValueError):
        normalized_header_row = 1

    lower_name = str(file_name or "").lower()
    if lower_name.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        result = _preview_xlsx(
            file_bytes,
            requested_sheet=requested_sheet,
            header_row=normalized_header_row,
        )
    elif lower_name.endswith(".csv"):
        result = _preview_csv(file_bytes, header_row=normalized_header_row)
    else:
        raise HTTPException(
            status_code=400,
            detail="Серверный preview поддерживает .xlsx/.xlsm и .csv.",
        )

    return {
        "fileName": file_name,
        **result,
    }
