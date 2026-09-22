from __future__ import annotations

import io

from openpyxl import Workbook

import core.parsers as parsers
from backend.app.services.source_preview import preview_source_file
from core.parsers import load_file_polars


def _xlsx_with_custom_sheet_and_header() -> bytes:
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Cover"
    cover.append(["Отчёт за период"])
    cover.append(["Это не таблица"])

    data = workbook.create_sheet("Transactions")
    data.append(["ООО Example"])
    data.append(["Сверка за август"])
    data.append(["Дата", "RRN", "Сумма"])
    data.append(["31.08.2026", "00123", 100000])
    data.append(["01.09.2026", "00124", 200000])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_excel_named_sheet_and_header_row_are_applied():
    frame = load_file_polars(
        _xlsx_with_custom_sheet_and_header(),
        "sample.xlsx",
        sheet_name="Transactions",
        header_row=3,
    )

    assert frame.columns == ["Дата", "RRN", "Сумма"]
    assert frame.height == 2
    assert str(frame["RRN"][0]) == "00123"
    assert frame["Сумма"][1] == 200000


def test_named_sheet_custom_header_uses_calamine_before_pandas(monkeypatch):
    def fail_pandas(*_args, **_kwargs):
        raise AssertionError("pandas.read_excel fallback should not be used")

    monkeypatch.setattr(parsers.pd, "read_excel", fail_pandas)

    frame = load_file_polars(
        _xlsx_with_custom_sheet_and_header(),
        "sample.xlsx",
        sheet_name="Transactions",
        header_row=3,
    )

    assert frame.columns == ["Дата", "RRN", "Сумма"]
    assert frame.height == 2
    assert str(frame["RRN"][0]) == "00123"


def test_csv_header_row_is_applied_before_delimiter_detection():
    content = (
        "Отчёт эквайринга\n"
        "Период;Август 2026\n"
        "Дата;RRN;Сумма\n"
        "31.08.2026;0001;100\n"
        "01.09.2026;0002;200\n"
    ).encode("utf-8")

    frame = load_file_polars(
        content,
        "sample.csv",
        header_row=3,
    )

    assert frame.columns == ["Дата", "RRN", "Сумма"]
    assert frame.height == 2
    assert frame["Дата"][0] == "31.08.2026"


def test_server_preview_reads_named_sheet_and_real_row_count():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результат запроса."
    sheet.append(["Дата", "RRN", "Summa"])
    sheet.append(["01.09.2026", "030793201824", 1000])
    sheet.append(["02.09.2026", "030793201825", 2000])

    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = preview_source_file(
        buffer.getvalue(),
        "66000107 данные алифа.xlsx",
        header_row=1,
    )

    assert preview["selectedSheet"] == "Результат запроса."
    assert preview["sheetNames"] == ["Результат запроса."]
    assert preview["columns"] == ["Дата", "RRN", "Summa"]
    assert preview["rowCount"] == 2
    assert preview["rows"][0]["RRN"] == "030793201824"
    assert preview["parser"] == "server-openpyxl"


def test_server_preview_matches_requested_sheet_after_trim_and_casefold():
    workbook = Workbook()
    first = workbook.active
    first.title = "Cover"
    first.append(["ignore"])

    data = workbook.create_sheet("Результат запроса.")
    data.append(["Дата", "RRN", "Summa"])
    data.append(["01.09.2026", "030793201824", 1000])

    buffer = io.BytesIO()
    workbook.save(buffer)

    preview = preview_source_file(
        buffer.getvalue(),
        "sample.xlsx",
        requested_sheet="  результат запроса.  ",
        header_row=1,
    )

    assert preview["selectedSheet"] == "Результат запроса."
    assert preview["columns"] == ["Дата", "RRN", "Summa"]

